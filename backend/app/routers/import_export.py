"""Import/Export router: Excel import/export for patents, trademarks, fees, old bonus."""
import io
import json
from datetime import datetime, date, timedelta
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_
from app.database import get_db
from app.models import (
    Patent, Trademark, PatentFee, Agency, AuditLog, SysUser, Employee,
    PatentInventor, BonusRule, PatentBonusBatch, PatentBonusItem,
    PatentBonusDetail, FeeStandard
)
from app.security import get_current_user, get_current_user_flexible
from app import fee_calculator
from app.patent_match import match_patent_by_name
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

router = APIRouter(prefix="/api/io", tags=["import_export"])


# ===== Export helpers =====
def _style_header(ws, col_count):
    """Apply styling to header row."""
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="409EFF", end_color="409EFF", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border


# ===== Patent Export =====
@router.get("/export/patents")
def export_patents(
    search: str = Query(""),
    patent_type: str = Query(""),
    status: str = Query(""),
    db: Session = Depends(get_db),
    current_user: SysUser = Depends(get_current_user_flexible)
):
    """Export patents to Excel."""
    q = db.query(Patent).filter(Patent.is_deleted == False)
    if search:
        like = f"%{search}%"
        q = q.filter(Patent.patent_name.like(like) | Patent.application_no.like(like) | Patent.inventors.like(like))
    if patent_type:
        q = q.filter(Patent.patent_type == patent_type)
    if status:
        q = q.filter(Patent.status == status)

    patents = q.order_by(Patent.application_date.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "专利清单"

    headers = [
        "申请号", "名称", "类型", "状态", "申请日", "授权号", "授权日", "首次公开日",
        "申请人", "发明人", "IPC分类号", "对应项目", "对应产品", "保护要素",
        "快速预审", "优审", "费减", "费减比例", "官费", "代理费", "预审费",
        "第1年年费", "第2年年费", "第3年年费", "第4年年费", "第5年年费",
        "第6年年费", "第7年年费", "第8年年费", "第9年年费", "第10年年费",
        "PCT国际专利", "代理案件号"
    ]
    ws.append(headers)
    _style_header(ws, len(headers))

    for p in patents:
        row = [
            p.application_no, p.patent_name, p.patent_type, p.status,
            str(p.application_date) if p.application_date else "",
            p.authorization_no or "",
            str(p.authorization_date) if p.authorization_date else "",
            str(p.first_publication_date) if p.first_publication_date else "",
            p.applicant, p.inventors or "",
            p.ipc_classification or "", p.correspondence_project or "",
            p.correspondence_product or "", p.protection_element or "",
            "是" if p.quick_examination else "否",
            "是" if p.expedited_examination else "否",
            "是" if p.fee_reduction else "否",
            p.fee_reduction_rate or "",
            p.official_fee or "", p.agency_fee or "", p.pre_examination_fee or "",
            p.fee_year1 or "", p.fee_year2 or "", p.fee_year3 or "",
            p.fee_year4 or "", p.fee_year5 or "", p.fee_year6 or "",
            p.fee_year7 or "", p.fee_year8 or "", p.fee_year9 or "", p.fee_year10 or "",
            "是" if p.is_pct else "否",
            p.agency_case_no or ""
        ]
        ws.append(row)

    # Auto-width
    for col_idx, header in enumerate(headers, 1):
        max_len = max(len(str(header)), max((len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)), default=0))
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Audit log
    log = AuditLog(user_id=current_user.id, user_name=current_user.display_name,
                   action="export", business_type="patent")
    db.add(log)
    db.commit()

    filename = f"专利清单_{date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    )


# ===== Trademark Export =====
@router.get("/export/trademarks")
def export_trademarks(
    search: str = Query(""),
    status: str = Query(""),
    scope_group: str = Query(""),
    db: Session = Depends(get_db),
    current_user: SysUser = Depends(get_current_user_flexible)
):
    """Export trademarks to Excel."""
    q = db.query(Trademark).filter(Trademark.is_deleted == False)
    if search:
        like = f"%{search}%"
        q = q.filter(Trademark.trademark_name.like(like) | Trademark.trademark_no.like(like))
    if status:
        q = q.filter(Trademark.status == status)
    if scope_group:
        q = q.filter(Trademark.scope_group == scope_group)

    trademarks = q.order_by(Trademark.application_date.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "商标清单"

    headers = [
        "商标注册号", "商标名称", "商标类型", "申请日", "注册日", "有效期至",
        "尼斯分类", "注册范围组", "申请人", "状态", "已续展次数",
        "最近续展日", "续展截止日", "宽展期截止日", "部分授权类别", "备注"
    ]
    ws.append(headers)
    _style_header(ws, len(headers))

    for t in trademarks:
        ws.append([
            t.trademark_no or "", t.trademark_name, t.trademark_type or "",
            str(t.application_date) if t.application_date else "",
            str(t.registration_date) if t.registration_date else "",
            str(t.valid_until) if t.valid_until else "",
            t.nice_class or "", t.scope_group or "",
            t.applicant, t.status,
            t.renewal_count or 0,
            str(t.last_renewal_date) if t.last_renewal_date else "",
            str(t.renewal_deadline) if t.renewal_deadline else "",
            str(t.grace_period_deadline) if t.grace_period_deadline else "",
            t.partial_grant_class or "", t.notes or ""
        ])

    for col_idx, header in enumerate(headers, 1):
        max_len = max(len(str(header)), max((len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)), default=0))
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    log = AuditLog(user_id=current_user.id, user_name=current_user.display_name,
                   action="export", business_type="trademark")
    db.add(log)
    db.commit()

    filename = f"商标清单_{date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    )


# ===== Fee Export =====
@router.get("/export/fees")
def export_fees(
    status: str = Query(""),
    db: Session = Depends(get_db),
    current_user: SysUser = Depends(get_current_user_flexible)
):
    """Export patent fees to Excel."""
    q = db.query(PatentFee).join(Patent).filter(Patent.is_deleted == False)
    if status:
        q = q.filter(PatentFee.status == status)

    fees = q.order_by(PatentFee.due_date.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "专利年费"

    headers = [
        "专利名称", "申请号", "专利类型", "年度", "应缴日期", "实缴日期",
        "标准金额", "费减金额", "实付金额", "状态", "滞纳金", "收据号", "缴费来源"
    ]
    ws.append(headers)
    _style_header(ws, len(headers))

    for f in fees:
        p = f.patent
        ws.append([
            p.patent_name if p else "", p.application_no if p else "",
            p.patent_type if p else "", f.fee_year,
            str(f.due_date) if f.due_date else "",
            str(f.actual_pay_date) if f.actual_pay_date else "",
            f.standard_amount or "", f.reduced_amount or "",
            f.actual_pay_amount or "", f.status,
            f.late_fee or "", f.receipt_no or "",
            f.payment_source or ""
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"专利年费_{date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    )


# ===== Bonus Export =====
@router.get("/export/bonus/{batch_id}")
def export_bonus(
    batch_id: int,
    db: Session = Depends(get_db),
    current_user: SysUser = Depends(get_current_user_flexible)
):
    """Export bonus batch details to Excel."""
    batch = db.query(PatentBonusBatch).filter(PatentBonusBatch.id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="批次不存在")

    items = db.query(PatentBonusItem).filter(PatentBonusItem.batch_id == batch_id).all()

    wb = openpyxl.Workbook()

    # Sheet 1: Summary
    ws1 = wb.active
    ws1.title = "奖金汇总"
    ws1.append(["批次编号", "批次名称", "总金额", "条目数", "状态", "周期开始", "周期结束"])
    _style_header(ws1, 7)
    ws1.append([
        batch.batch_code, batch.batch_name, batch.total_amount or 0,
        len(items), batch.status,
        str(batch.period_start) if batch.period_start else "",
        str(batch.period_end) if batch.period_end else ""
    ])

    # Sheet 2: Items
    ws2 = wb.create_sheet("奖金明细")
    ws2.append([
        "专利名称", "专利类型", "里程碑", "研发部门",
        "应发奖金", "已发奖金", "本次申请发放",
        "基础金额", "PCT叠加", "总金额",
        "发放状态", "发放日期", "发明人", "贡献比例", "分配金额", "发放状态(个人)"
    ])
    _style_header(ws2, 16)

    for item in items:
        details = db.query(PatentBonusDetail).filter(PatentBonusDetail.bonus_item_id == item.id).all()
        if not details:
            ws2.append([
                item.patent_name, item.patent_type, item.milestone,
                item.department or "",
                item.approved_amount or 0, item.received_amount or 0, item.applied_amount or 0,
                item.base_amount, item.pct_bonus, item.total_amount,
                item.payment_status,
                str(item.payment_date) if item.payment_date else "",
                "", "", "", ""
            ])
        else:
            for d in details:
                inventor_name = ""
                if d.inventor_id:
                    emp = db.query(Employee).filter(Employee.id == d.inventor_id).first()
                    inventor_name = emp.name if emp else ""
                elif d.external_name:
                    inventor_name = d.external_name

                ws2.append([
                    item.patent_name, item.patent_type, item.milestone,
                    item.department or "",
                    item.approved_amount or 0, item.received_amount or 0, item.applied_amount or 0,
                    item.base_amount, item.pct_bonus, item.total_amount,
                    item.payment_status,
                    str(item.payment_date) if item.payment_date else "",
                    inventor_name, d.contribution_ratio, d.amount, d.payment_status
                ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"奖金明细_{batch.batch_code}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    )


# ===== Fee Standard Export =====
@router.get("/export/fee-standards")
def export_fee_standards(
    db: Session = Depends(get_db),
    current_user: SysUser = Depends(get_current_user_flexible)
):
    """Export fee standards to Excel."""
    standards = db.query(FeeStandard).order_by(FeeStandard.category, FeeStandard.fee_type).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "费用标准"

    headers = ["类别", "费用类型", "专利类型", "年份起", "年份止", "标准金额", "费减85%", "费减70%", "单位", "备注"]
    ws.append(headers)
    _style_header(ws, len(headers))

    for s in standards:
        ws.append([
            s.category or "", s.fee_type, s.patent_type or "",
            s.year_start or "", s.year_end or "",
            s.standard_amount or "", s.reduced_85_amount or "",
            s.reduced_70_amount or "", s.unit or "", s.note or ""
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"费用标准_{date.today().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    )


# ===== Patent Import =====
@router.post("/import/patents")
async def import_patents(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: SysUser = Depends(get_current_user)
):
    """Import patents from Excel. Supports both system export template and the user's 专利清单 template."""
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    def parse_date(val):
        if not val:
            return None
        if isinstance(val, (date, datetime)):
            return val if isinstance(val, date) else val.date()
        for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d"):
            try:
                return datetime.strptime(str(val).strip(), fmt).date()
            except:
                continue
        return None

    def parse_bool(val):
        if not val:
            return False
        return str(val).strip() in ("是", "True", "true", "1", "YES", "yes", "Y", "y", "✓")

    def parse_float(val):
        if val is None or val == "":
            return None
        try:
            return float(val)
        except:
            return None

    # Header mapping: supports Chinese headers from 专利清单.xlsx and system export headers
    HEADER_MAP = {
        # 专利清单 template
        "申请号": "application_no",
        "专利名称": "patent_name",
        "名称": "patent_name",
        "申请类型": "patent_type",
        "类型": "patent_type",
        "法律状态": "status",
        "状态": "status",
        "申请日期": "application_date",
        "申请日": "application_date",
        "授权号": "authorization_no",
        "授权日期": "authorization_date",
        "授权日": "authorization_date",
        "首次公开日": "first_publication_date",
        "发明人": "inventors",
        "权利人": "applicant",
        "申请人": "applicant",
        "代理事务所": "agency_name",
        "代理机构": "agency_name",
        "对应项目": "correspondence_project",
        "对应产品": "correspondence_product",
        "保护要素": "protection_element",
        "快速预审": "quick_examination",
        "优审": "expedited_examination",
        "费减": "fee_reduction",
        "费减比例": "fee_reduction_rate",
        "官费": "official_fee",
        "代理费": "agency_fee",
        "预审费": "pre_examination_fee",
        "第一年年费": "fee_year1",
        "第二年年费": "fee_year2",
        "第三年年费": "fee_year3",
        "第四年年费": "fee_year4",
        "第五年年费": "fee_year5",
        "第六年年费": "fee_year6",
        "第七年年费": "fee_year7",
        "第八年年费": "fee_year8",
        "第九年年费": "fee_year9",
        "第十年年费": "fee_year10",
        "第1年年费": "fee_year1",
        "第2年年费": "fee_year2",
        "第3年年费": "fee_year3",
        "第4年年费": "fee_year4",
        "第5年年费": "fee_year5",
        "第6年年费": "fee_year6",
        "第7年年费": "fee_year7",
        "第8年年费": "fee_year8",
        "第9年年费": "fee_year9",
        "第10年年费": "fee_year10",
    }

    # Find header row
    header_row_idx = None
    headers = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        cells = [str(v).strip() if v is not None else "" for v in row]
        mapped = [HEADER_MAP.get(c) for c in cells]
        if any(mapped):
            header_row_idx = idx
            headers = cells
            break

    if header_row_idx is None:
        raise HTTPException(status_code=400, detail="未找到有效的表头行")

    col_map = {}
    for i, h in enumerate(headers):
        if h in HEADER_MAP:
            col_map[HEADER_MAP[h]] = i

    if "application_no" not in col_map:
        raise HTTPException(status_code=400, detail="表头缺少申请号列")

    success, failed, skipped = 0, 0, 0
    rows = list(ws.iter_rows(min_row=header_row_idx + 1, values_only=True))

    # Cache agencies by name
    agencies = {a.agency_name: a.id for a in db.query(Agency).filter(Agency.is_deleted == False).all()}

    for row in rows:
        try:
            app_no = str(row[col_map["application_no"]]).strip() if row[col_map["application_no"]] else ""
            if not app_no:
                skipped += 1
                continue

            existing = db.query(Patent).filter(Patent.application_no == app_no, Patent.is_deleted == False).first()
            if existing:
                skipped += 1
                continue

            def get(col, default=None):
                if col not in col_map:
                    return default
                return row[col_map[col]]

            patent_type_raw = str(get("patent_type", "") or "").strip()
            type_map = {
                "发明": "发明", "invention": "发明", "發明": "发明",
                "实用新型": "实用新型", "utility model": "实用新型",
                "外观": "外观", "design": "外观",
                "软著": "软著", "软件著作权": "软著", "software copyright": "软著",
                "软产": "软产"
            }
            patent_type = type_map.get(patent_type_raw.lower(), patent_type_raw or "发明")

            status_raw = str(get("status", "") or "").strip()
            status_map = {
                "受理": "受理", "实质审查": "实质审查", "授权": "授权", "已登记": "已登记",
                "驳回": "驳回", "复审": "复审", "放弃": "放弃", "申请中": "申请中",
                "issued": "授权", "granted": "授权"
            }
            status = status_map.get(status_raw.lower(), status_raw or "申请中")

            agency_name = get("agency_name")
            agency_id = agencies.get(agency_name) if agency_name else None

            p = Patent(
                application_no=app_no,
                patent_name=str(get("patent_name") or "").strip() or app_no,
                patent_type=patent_type,
                status=status,
                application_date=parse_date(get("application_date")),
                authorization_no=str(get("authorization_no")).strip() if get("authorization_no") else None,
                authorization_date=parse_date(get("authorization_date")),
                first_publication_date=parse_date(get("first_publication_date")),
                applicant=str(get("applicant") or "示例科技有限公司").strip(),
                inventors=str(get("inventors")).strip() if get("inventors") else None,
                correspondence_project=str(get("correspondence_project")).strip() if get("correspondence_project") else None,
                correspondence_product=str(get("correspondence_product")).strip() if get("correspondence_product") else None,
                protection_element=str(get("protection_element")).strip() if get("protection_element") else None,
                quick_examination=parse_bool(get("quick_examination")),
                expedited_examination=parse_bool(get("expedited_examination")),
                fee_reduction=parse_bool(get("fee_reduction")),
                fee_reduction_rate=str(get("fee_reduction_rate")).strip() if get("fee_reduction_rate") else None,
                official_fee=parse_float(get("official_fee")),
                agency_fee=parse_float(get("agency_fee")),
                pre_examination_fee=parse_float(get("pre_examination_fee")),
                fee_year1=parse_float(get("fee_year1")),
                fee_year2=parse_float(get("fee_year2")),
                fee_year3=parse_float(get("fee_year3")),
                fee_year4=parse_float(get("fee_year4")),
                fee_year5=parse_float(get("fee_year5")),
                fee_year6=parse_float(get("fee_year6")),
                fee_year7=parse_float(get("fee_year7")),
                fee_year8=parse_float(get("fee_year8")),
                fee_year9=parse_float(get("fee_year9")),
                fee_year10=parse_float(get("fee_year10")),
                agency_id=agency_id,
                created_by=current_user.id
            )
            db.add(p)
            db.flush()  # 拿到 p.id

            # 自动生成年费记录（应缴金额按规则算）
            fees = fee_calculator.generate_patent_fees(db, p, commit=False)
            fee_map = {f.fee_year: f for f in fees}

            # 回填 Excel 中的实际已缴年费金额（含滞纳金 "135+5" 格式）
            for i in range(1, 11):
                raw = get(f"fee_year{i}")
                if raw is None:
                    continue
                s = str(raw).strip()
                if s in ("", "/", "—", "-"):
                    continue  # 免缴 / 未填
                fee = fee_map.get(i)
                if fee is None:
                    continue
                if "+" in s:
                    parts = s.split("+")
                    try:
                        fee.actual_pay_amount = float(parts[0])
                        fee.late_fee = float(parts[1]) if len(parts) > 1 else 0
                    except ValueError:
                        continue
                else:
                    try:
                        fee.actual_pay_amount = float(s)
                    except ValueError:
                        continue
                fee.status = "已缴"

            success += 1
        except Exception as e:
            failed += 1
            continue

    db.commit()

    log = AuditLog(user_id=current_user.id, user_name=current_user.display_name,
                   action="import", business_type="patent",
                   after_value=json.dumps({"file": file.filename, "success": success, "failed": failed, "skipped": skipped}))
    db.add(log)
    db.commit()

    return {"message": f"导入完成：成功 {success} 条，失败 {failed} 条，跳过 {skipped} 条", "success": success, "failed": failed, "skipped": skipped}


# ===== Trademark Import =====
@router.post("/import/trademarks")
async def import_trademarks(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: SysUser = Depends(get_current_user)
):
    """Import trademarks from Excel."""
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    success, failed, skipped = 0, 0, 0

    for row in rows:
        if not row or not row[1]:
            skipped += 1
            continue
        try:
            tm_name = str(row[1]).strip()
            if not tm_name:
                skipped += 1
                continue

            def parse_date(val):
                if not val:
                    return None
                if isinstance(val, (date, datetime)):
                    return val if isinstance(val, date) else val.date()
                try:
                    return datetime.strptime(str(val).strip(), "%Y-%m-%d").date()
                except:
                    return None

            tm_no = str(row[0]).strip() if row[0] else None
            # Check duplicate
            if tm_no:
                existing = db.query(Trademark).filter(Trademark.trademark_no == tm_no, Trademark.is_deleted == False).first()
                if existing:
                    skipped += 1
                    continue

            valid_until = parse_date(row[5])
            if not valid_until:
                skipped += 1
                continue

            t = Trademark(
                trademark_no=tm_no,
                trademark_name=tm_name,
                trademark_type=str(row[2]).strip() if row[2] else None,
                application_date=parse_date(row[3]),
                registration_date=parse_date(row[4]),
                valid_until=valid_until,
                nice_class=str(row[6]).strip() if row[6] else None,
                scope_group=str(row[7]).strip() if row[7] else None,
                applicant=str(row[8]).strip() if row[8] else "示例科技有限公司",
                status=str(row[9]).strip() if row[9] else "申请中",
                renewal_count=int(row[10]) if row[10] else 0,
                last_renewal_date=parse_date(row[11]),
                renewal_deadline=parse_date(row[12]),
                grace_period_deadline=parse_date(row[13]),
                partial_grant_class=str(row[14]).strip() if row[14] else None,
                notes=str(row[15]).strip() if row[15] else None,
                created_by=current_user.id
            )
            db.add(t)
            success += 1
        except Exception as e:
            failed += 1
            continue

    db.commit()

    log = AuditLog(user_id=current_user.id, user_name=current_user.display_name,
                   action="import", business_type="trademark",
                   after_value=json.dumps({"file": file.filename, "success": success, "failed": failed, "skipped": skipped}))
    db.add(log)
    db.commit()

    return {"message": f"导入完成：成功 {success} 条，失败 {failed} 条，跳过 {skipped} 条", "success": success, "failed": failed, "skipped": skipped}


# ===== Fee Import (from agency Excel) =====
@router.post("/import/fees")
async def import_fees(
    file: UploadFile = File(...),
    agency_id: int = Query(None),
    db: Session = Depends(get_db),
    current_user: SysUser = Depends(get_current_user)
):
    """Import patent fees from agency Excel.
    Expected columns: 申请号/专利号, 年度, 应缴日期, 实付金额, 缴费日期, 收据号
    """
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    success, failed, skipped = 0, 0, 0
    deviations = []

    for row in rows:
        if not row or not row[0]:
            skipped += 1
            continue
        try:
            app_no = str(row[0]).strip()
            patent = db.query(Patent).filter(
                Patent.application_no == app_no, Patent.is_deleted == False
            ).first()

            if not patent:
                skipped += 1
                continue

            fee_year = int(row[1]) if row[1] else None
            if not fee_year:
                skipped += 1
                continue

            def parse_date(val):
                if not val:
                    return None
                if isinstance(val, (date, datetime)):
                    return val if isinstance(val, date) else val.date()
                try:
                    return datetime.strptime(str(val).strip(), "%Y-%m-%d").date()
                except:
                    return None

            def parse_float(val):
                if val is None or val == "":
                    return None
                try:
                    return float(val)
                except:
                    return None

            actual_amount = parse_float(row[3]) if len(row) > 3 else None

            # Find standard amount from fee_standard
            standard_amount = None
            ptype = patent.patent_type
            std = db.query(FeeStandard).filter(
                FeeStandard.fee_type == "年费",
                FeeStandard.patent_type == ptype
            ).filter(
                FeeStandard.year_start <= fee_year,
                FeeStandard.year_end >= fee_year
            ).first()
            if std:
                if patent.fee_reduction and patent.fee_reduction_rate == "85%":
                    standard_amount = std.reduced_85_amount
                elif patent.fee_reduction and patent.fee_reduction_rate == "70%":
                    standard_amount = std.reduced_70_amount
                else:
                    standard_amount = std.standard_amount

            # Check deviation
            if actual_amount and standard_amount and standard_amount > 0:
                dev = abs(actual_amount - standard_amount) / standard_amount
                if dev > 0.05:
                    deviations.append({
                        "patent_name": patent.patent_name,
                        "fee_year": fee_year,
                        "standard": standard_amount,
                        "actual": actual_amount,
                        "deviation": actual_amount - standard_amount
                    })

            # Check if fee record already exists
            existing = db.query(PatentFee).filter(
                PatentFee.patent_id == patent.id,
                PatentFee.fee_year == fee_year
            ).first()

            if existing:
                # Update
                if actual_amount:
                    existing.actual_pay_amount = actual_amount
                    existing.actual_pay_date = parse_date(row[4]) if len(row) > 4 else None
                    existing.receipt_no = str(row[5]).strip() if len(row) > 5 and row[5] else None
                    existing.status = "已缴"
                    existing.agency_id = agency_id or existing.agency_id
                    existing.imported_from = file.filename
                skipped += 1
            else:
                fee = PatentFee(
                    patent_id=patent.id,
                    fee_year=fee_year,
                    due_date=parse_date(row[2]) if len(row) > 2 else None,
                    actual_pay_amount=actual_amount,
                    actual_pay_date=parse_date(row[4]) if len(row) > 4 else None,
                    standard_amount=standard_amount,
                    receipt_no=str(row[5]).strip() if len(row) > 5 and row[5] else None,
                    status="已缴" if actual_amount else "待缴",
                    agency_id=agency_id,
                    payment_source="代理代缴" if agency_id else "官费",
                    imported_from=file.filename
                )
                db.add(fee)
                success += 1
        except Exception as e:
            failed += 1
            continue

    db.commit()

    log = AuditLog(user_id=current_user.id, user_name=current_user.display_name,
                   action="import", business_type="patent_fee",
                   after_value=json.dumps({"file": file.filename, "success": success, "failed": failed, "skipped": skipped, "deviations": len(deviations)}))
    db.add(log)
    db.commit()

    return {
        "message": f"导入完成：成功 {success} 条，更新 {skipped} 条，失败 {failed} 条",
        "success": success, "failed": failed, "skipped": skipped,
        "deviations": deviations
    }


# ===== Old Bonus Import (supports English header format) =====
@router.post("/import/old-bonus")
async def import_old_bonus(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: SysUser = Depends(get_current_user)
):
    """Import bonus records from Excel.
    Supports English header format from user's yearly bonus tables:
    No./Name/Type/Applied Number/Applied date/Issued date/R&D department/
    The bonus approved/Received bonus/Applied bonus of this time/Inventor/Status
    Also supports old Chinese format: 序号/申请号/名称/类型/发明人/奖金金额/发放日期/备注
    """
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    def parse_date(val):
        if not val:
            return None
        if isinstance(val, (date, datetime)):
            return val if isinstance(val, date) else val.date()
        for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d"):
            try:
                return datetime.strptime(str(val).strip(), fmt).date()
            except:
                continue
        return None

    def parse_float(val):
        if val is None or val == "":
            return 0
        try:
            return float(val)
        except:
            return 0

    # Header mapping: English headers -> field names, Chinese headers -> field names
    BONUS_HEADER_MAP = {
        # English headers (user's format - 2022+)
        "no.": "seq_no",
        "name": "patent_name",
        "type": "patent_type",
        "applied number": "application_no",
        "applied date": "applied_date",
        "issued date": "issued_date",
        "r&d department": "department",
        "the bonus approved": "approved_amount",
        "received bonus": "received_amount",
        "applied bonus of this time": "applied_amount",
        "inventor": "inventor_str",
        "status": "bonus_status",
        # Legacy inventor-based format (2020/2021)
        "bonus amount (cny)": "legacy_amount",
        "bonus amount": "legacy_amount",
        "total": "legacy_amount",
        "patents": "legacy_patents",
        # Chinese headers (old format)
        "序号": "seq_no",
        "申请号": "application_no",
        "名称": "patent_name",
        "专利名称": "patent_name",
        "类型": "patent_type",
        "申请类型": "patent_type",
        "发明人": "inventor_str",
        "奖金金额": "applied_amount",
        "发放日期": "payment_date",
        "备注": "remark",
        "研发部门": "department",
        "应发奖金": "approved_amount",
        "已发奖金": "received_amount",
        "本次申请": "applied_amount",
    }

    # Limit columns to avoid openpyxl phantom columns (xlsx may report 16384 cols)
    MAX_SCAN_COLS = 30

    # Find header row (skip title rows like "2025 IP bonus list")
    header_row_idx = None
    headers = []
    for idx, row in enumerate(ws.iter_rows(values_only=True, max_col=MAX_SCAN_COLS), 1):
        cells = [str(v).strip().lower() if v is not None else "" for v in row]
        mapped = [BONUS_HEADER_MAP.get(c) for c in cells]
        if sum(1 for m in mapped if m) >= 2:  # at least 2 known headers (legacy format has fewer)
            header_row_idx = idx
            headers = cells
            break

    if header_row_idx is None:
        raise HTTPException(status_code=400, detail="未找到有效的表头行，请确保Excel包含 No./Name/Type 等列名")

    # Build column map
    col_map = {}
    for i, h in enumerate(headers):
        field = BONUS_HEADER_MAP.get(h)
        if field and field not in col_map:
            col_map[field] = i

    # Detect legacy inventor-based format (2020/2021 files: Name + Bonus Amount, no Applied Number/Type)
    is_legacy = "legacy_amount" in col_map and "application_no" not in col_map
    if is_legacy and "patent_name" in col_map:
        # In legacy format, "Name" column is the inventor, not the patent name
        col_map["inventor_str"] = col_map.pop("patent_name")

    # Extract year (and optionally month) from filename for batch naming
    import re
    date_match = re.search(r'20(\d{2})(\d{2})?', file.filename or "")
    batch_year = int("20" + date_match.group(1)) if date_match else datetime.now().year
    batch_month = int(date_match.group(2)) if date_match and date_match.group(2) else None

    # Find or create "旧规则-导入" rule
    rule = db.query(BonusRule).filter(BonusRule.rule_name.like("%旧规则%")).first()
    if not rule:
        rule = BonusRule(
            rule_name="旧规则-导入",
            effective_from=date(2000, 1, 1),
            effective_until=date(2030, 12, 31),
            is_current=False,
            rules_json="{}",
            raw_text="从Excel导入的历史奖金记录"
        )
        db.add(rule)
        db.commit()
        db.refresh(rule)

    # Build batch name and code (include month when available for same-year multi-batch)
    if batch_month:
        batch_name = f"{batch_year}年{batch_month:02d}月奖金发放"
        batch_code = f"IMPORT-{batch_year}{batch_month:02d}"
    else:
        batch_name = f"{batch_year}年度奖金发放"
        batch_code = f"IMPORT-{batch_year}"

    existing_batch = db.query(PatentBonusBatch).filter(
        PatentBonusBatch.batch_code == batch_code
    ).first()

    # Compute batch period
    if batch_month:
        period_start = date(batch_year, batch_month, 1)
        period_end = date(batch_year, 12, 31) if batch_month == 12 else date(batch_year, batch_month + 1, 1) - timedelta(days=1)
    else:
        period_start = date(batch_year, 1, 1)
        period_end = date(batch_year, 12, 31)

    if existing_batch:
        # Delete existing details first (FK constraint), then items
        old_items = db.query(PatentBonusItem).filter(
            PatentBonusItem.batch_id == existing_batch.id
        ).all()
        old_item_ids = [item.id for item in old_items]
        if old_item_ids:
            db.query(PatentBonusDetail).filter(
                PatentBonusDetail.bonus_item_id.in_(old_item_ids)
            ).delete(synchronize_session=False)
        db.query(PatentBonusItem).filter(
            PatentBonusItem.batch_id == existing_batch.id
        ).delete(synchronize_session=False)
        db.commit()
        batch = existing_batch
        batch.batch_name = batch_name
        batch.period_start = period_start
        batch.period_end = period_end
    else:
        batch = PatentBonusBatch(
            batch_code=batch_code,
            batch_name=batch_name,
            rule_id=rule.id,
            period_start=period_start,
            period_end=period_end,
            status="已发放",
            created_by=current_user.id
        )
        db.add(batch)
        db.commit()
        db.refresh(batch)

    rows = list(ws.iter_rows(min_row=header_row_idx + 1, values_only=True, max_col=MAX_SCAN_COLS))
    success, failed = 0, 0
    total_amount = 0.0
    total_approved = 0.0
    total_received = 0.0

    # Type mapping
    def map_type(raw):
        if not raw:
            return "发明"
        raw_lower = str(raw).strip().lower()
        type_map = {
            "invention": "发明", "发明": "发明", "發明": "发明",
            "utility model": "实用新型", "实用新型": "实用新型",
            "design": "外观", "外观": "外观",
            "software copyright": "软著", "软著": "软著", "软件著作权": "软著",
            "软产": "软产"
        }
        return type_map.get(raw_lower, str(raw).strip())

    # Status mapping
    def map_status(raw):
        if not raw:
            return "已发"
        raw_str = str(raw).strip().lower()
        if any(k in raw_str for k in ["paid", "已发", "已付", "完成"]):
            return "已发"
        if any(k in raw_str for k in ["pending", "待发", "未发", "申请中"]):
            return "待发"
        return "已发"

    for row in rows:
        if not row or not any(v is not None and str(v).strip() for v in row):
            continue

        # Skip summary/total rows
        first_cell = str(row[0]).strip().lower() if row[0] else ""
        if any(k in first_cell for k in ["total", "合计", "sum", "total amount"]):
            continue

        try:
            def get(field, default=None):
                if field not in col_map:
                    return default
                idx = col_map[field]
                val = row[idx] if idx < len(row) else None
                return val

            # === Legacy inventor-based format (2020/2021 files) ===
            if is_legacy:
                inventor_name = str(get("inventor_str") or "").strip()
                if not inventor_name:
                    continue
                amount = parse_float(get("legacy_amount"))
                if amount <= 0:
                    continue
                patent_names_str = str(get("legacy_patents") or "").strip()
                patent_name = patent_names_str if patent_names_str else f"历史奖金-{batch_year}"
                # 按名称匹配已有专利（legacy 条目常带中文/英文混合格式，能匹配就关联，聚合条目保持 NULL）
                legacy_patent = None
                if patent_names_str:
                    legacy_patent = match_patent_by_name(db, patent_names_str)

                item = PatentBonusItem(
                    batch_id=batch.id,
                    patent_id=legacy_patent.id if legacy_patent else None,
                    patent_type="发明",
                    patent_name=patent_name,
                    milestone="旧规则",
                    department=None,
                    approved_amount=amount,
                    received_amount=amount,
                    applied_amount=amount,
                    rule_id=rule.id,
                    base_amount=amount,
                    pct_bonus=0,
                    total_amount=amount,
                    payment_status="已发",
                    payment_date=None,
                    source="Excel导入"
                )
                db.add(item)
                db.flush()

                emp = db.query(Employee).filter(
                    Employee.name == inventor_name, Employee.is_deleted == False
                ).first()
                detail = PatentBonusDetail(
                    bonus_item_id=item.id,
                    inventor_id=emp.id if emp else None,
                    external_name=inventor_name if not emp else None,
                    contribution_ratio=100.0,
                    amount=amount,
                    payment_status="已发"
                )
                db.add(detail)

                total_amount += amount
                total_approved += amount
                total_received += amount
                success += 1
                continue

            # === Standard patent-based format (2022+ files) ===
            patent_name = str(get("patent_name") or "").strip()
            if not patent_name:
                continue

            app_no = str(get("application_no") or "").strip()
            patent = None
            if app_no:
                patent = db.query(Patent).filter(
                    Patent.application_no == app_no, Patent.is_deleted == False
                ).first()
            if not patent:
                # 申请号未命中时按名称三级匹配（括号中文/换行/包含），提升奖金-专利关联率
                patent = match_patent_by_name(db, patent_name)

            patent_type = map_type(get("patent_type"))

            approved_amount = parse_float(get("approved_amount"))
            received_amount = parse_float(get("received_amount"))
            applied_amount = parse_float(get("applied_amount"))

            # Use applied_amount as total if available, otherwise use approved_amount
            item_total = applied_amount if applied_amount > 0 else (approved_amount if approved_amount > 0 else 0)

            department = str(get("department") or "").strip() if get("department") else None
            inventor_str = str(get("inventor_str") or "").strip() if get("inventor_str") else ""

            issued_date = parse_date(get("issued_date"))
            payment_date = parse_date(get("payment_date")) if "payment_date" in col_map else issued_date
            bonus_status = map_status(get("bonus_status"))

            milestone = "授权" if issued_date else "旧规则"

            item = PatentBonusItem(
                batch_id=batch.id,
                patent_id=patent.id if patent else None,
                patent_type=patent_type,
                patent_name=patent_name,
                milestone=milestone,
                department=department,
                approved_amount=approved_amount,
                received_amount=received_amount,
                applied_amount=applied_amount,
                rule_id=rule.id,
                base_amount=item_total,
                pct_bonus=0,
                total_amount=item_total,
                payment_status=bonus_status,
                payment_date=payment_date or issued_date,
                source="Excel导入"
            )
            db.add(item)
            db.flush()

            # Create detail for each inventor
            if inventor_str:
                # Handle comma or semicolon separated names
                sep = "," if "," in inventor_str else (";" if ";" in inventor_str else "、")
                inventors = [i.strip() for i in inventor_str.split(sep) if i.strip()]
                if inventors:
                    ratio = round(100 / len(inventors), 2)
                    per_amount = round(item_total * ratio / 100, 2)
                    for inv_name in inventors:
                        emp = db.query(Employee).filter(
                            Employee.name == inv_name, Employee.is_deleted == False
                        ).first()
                        detail = PatentBonusDetail(
                            bonus_item_id=item.id,
                            inventor_id=emp.id if emp else None,
                            external_name=inv_name if not emp else None,
                            contribution_ratio=ratio,
                            amount=per_amount,
                            payment_status=bonus_status
                        )
                        db.add(detail)

            total_amount += item_total
            total_approved += approved_amount
            total_received += received_amount
            success += 1
        except Exception as e:
            failed += 1
            continue

    batch.total_amount = round(total_amount, 2)
    db.commit()

    log = AuditLog(user_id=current_user.id, user_name=current_user.display_name,
                   action="import", business_type="bonus",
                   after_value=json.dumps({"file": file.filename, "batch_id": batch.id, "year": batch_year, "success": success, "failed": failed, "total": total_amount}))
    db.add(log)
    db.commit()

    return {
        "message": f"导入完成：成功 {success} 条，失败 {failed} 条，申请发放金额 ¥{total_amount:,.2f}，应发 ¥{total_approved:,.2f}，已发 ¥{total_received:,.2f}",
        "batch_id": batch.id,
        "batch_code": batch_code,
        "batch_name": batch_name,
        "year": batch_year,
        "month": batch_month,
        "success": success,
        "failed": failed,
        "total_amount": round(total_amount, 2),
        "total_approved": round(total_approved, 2),
        "total_received": round(total_received, 2)
    }


# ===== Download Template =====
@router.get("/template/{template_type}")
def download_template(
    template_type: str,
    current_user: SysUser = Depends(get_current_user_flexible)
):
    """Download Excel template for import."""
    wb = openpyxl.Workbook()
    ws = wb.active

    if template_type == "patent":
        ws.title = "专利清单模板"
        headers = [
            "申请号", "名称", "类型", "状态", "申请日", "授权日", "首次公开日",
            "申请人", "发明人", "IPC分类号", "对应项目", "对应产品", "保护要素",
            "快速预审", "优审", "费减", "费减比例", "官费", "代理费", "预审费",
            "第1年年费", "第2年年费", "第3年年费", "第4年年费", "第5年年费",
            "第6年年费", "第7年年费", "第8年年费", "第9年年费", "第10年年费",
            "PCT国际专利", "代理案件号"
        ]
        ws.append(headers)
        ws.append([
            "202310344421.9", "示例专利名称", "发明", "授权", "2023-06-15", "2024-03-20", "",
            "示例科技有限公司", "张三,李四", "G06F", "项目A", "产品B", "保护要素描述",
            "否", "否", "是", "85%", 900, 5000, "",
            900, 900, 900, 1200, 1200, 1200, 2000, 2000, 2000, 4000,
            "否", "AGENCY-001"
        ])

    elif template_type == "trademark":
        ws.title = "商标清单模板"
        headers = [
            "商标注册号", "商标名称", "商标类型", "申请日", "注册日", "有效期至",
            "尼斯分类", "注册范围组", "申请人", "状态", "已续展次数",
            "最近续展日", "续展截止日", "宽展期截止日", "部分授权类别", "备注"
        ]
        ws.append(headers)
        ws.append([
            "12345678", "SensCore", "文字", "2020-03-15", "2021-03-15", "2031-03-14",
            "9", "通用设备", "示例科技有限公司", "已注册", 0,
            "", "", "", "", "核心品牌"
        ])

    elif template_type == "fee":
        ws.title = "年费导入模板"
        headers = ["申请号/专利号", "年度", "应缴日期", "实付金额", "缴费日期", "收据号"]
        ws.append(headers)
        ws.append(["202310344421.9", 3, "2026-06-15", 900, "2026-06-10", "R20260610001"])

    elif template_type == "old-bonus":
        ws.title = "年度奖金模板"
        headers = ["No.", "Name", "Type", "Applied Number", "Applied date", "Issued date",
                   "R&D department", "The bonus approved", "Received bonus",
                   "Applied bonus of this time", "Inventor", "Status"]
        ws.append(headers)
        ws.append([1, "示例专利名称", "Invention", "202310344421.9", "2023-06-15", "2024-03-20",
                   "研发一部", 8000, 4000, 4000, "张三,李四", "Paid"])

    else:
        raise HTTPException(status_code=400, detail="未知模板类型")

    _style_header(ws, len(headers))
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{template_type}_template.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    )
